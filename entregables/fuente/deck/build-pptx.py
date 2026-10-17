# Genera entregables/03-pitch-deck.pptx: deck corto de 8 slides de Sobra. (Identidad Cobalto), 16:9, editable.
# El texto que se dice va en las notas del orador de cada slide.
#
# Uso:
#   python entregables/fuente/deck/build-pptx.py [--font "Segoe UI"] [--html <archivo.html>]
import math
import sys
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt
from lxml import etree

HERE = Path(__file__).resolve().parent
IMG = HERE / "img"
OUT = HERE.parent.parent / "03-pitch-deck.pptx"
DEMO = "https://ideathon-agro-web.vercel.app/"

SANS = sys.argv[sys.argv.index("--font") + 1] if "--font" in sys.argv else "Segoe UI"
MONO = "Consolas"
HTML_OUT = Path(sys.argv[sys.argv.index("--html") + 1]).resolve() if "--html" in sys.argv else None

W, H = 1920, 1080
SAFE = 86  # 0,6 in
GAP = 16
MIN_BODY_PX = 40  # 20 pt
LH = 1.32  # alto de línea conservador para PowerPoint en Windows

# Paleta Sobra Cobalt Instrument
C = {
    "canvas": "eef3fd",
    "surface": "ffffff",
    "subtle_bg": "e6edfb",
    "line": "d3def3",
    "fg": "0a1433",
    "muted": "3d4d7a",
    "subtle": "5b6b94",
    "cobalt": "2150e8",
    "cobalt_dark": "0b1640",
    "cobalt_surface": "101f54",
    "cobalt_line": "1f2f66",
    "cobalt_light": "7aa7ff",
    "green": "17784a",
    "coral": "b5402a",
    "white": "ffffff",
}


def R(t, size, color="fg", b=False, mono=False, sp=None, link=None):
    if sp is None:
        sp = -0.02 * size if (b and not mono and size >= 64) else 0
    return dict(t=t, size=size, color=color, b=b, mono=mono, sp=sp, link=link)


def LBL(t, size=40, color="cobalt"):
    return R(t.upper(), size, color, mono=True, b=True, sp=0.04 * size)


def T(x, y, w, h, paras, align="left", anchor="top", line=1.05, role="content"):
    if isinstance(paras, dict):
        paras = [[paras]]
    elif paras and isinstance(paras[0], dict):
        paras = [paras]
    return dict(kind="text", x=x, y=y, w=w, h=h, paras=paras, align=align, anchor=anchor, line=line, role=role)


def I(name, x, y, w, h, pos=(0.5, 0.5), border=True, role="content", link=None):
    return dict(kind="img", name=name, x=x, y=y, w=w, h=h, pos=pos, border=border, role=role, link=link)


def B(x, y, w, h, fill, alpha=1.0, role="bg", border_color=None):
    return dict(kind="rect", x=x, y=y, w=w, h=h, fill=fill, alpha=alpha, role=role, border_color=border_color)


SLIDES = []


def slide(els, notes, dark=False, chrome=True):
    n = len(SLIDES) + 1
    if chrome:
        num_color = "cobalt_light" if dark else "subtle"
        els = els + [
            I("icon.png", 1760, 88, 64, 64, border=False, role="chrome"),
            T(1580, 936, 244, 54, R(f"{n:02d} / 08", 32, num_color, mono=True, b=True), align="right", role="chrome"),
        ]
    SLIDES.append(dict(els=els, notes=notes, dark=dark))


# 1 · Portada (Dark Cobalt)
slide([
    I("icon.png", 104, 104, 112, 112, border=False),
    T(240, 100, 560, 140, [R("Sobra", 96, "white", b=True), R(".", 96, "cobalt_light", b=True)]),
    T(104, 272, 740, 260, [
        [R("Del residuo", 88, "white", b=True)],
        [R("al mercado.", 88, "cobalt_light", b=True)],
    ], line=1.0),
    T(104, 556, 740, 200, R("El mapa que dice si tu biomasa cierra después del flete.", 44, "line"), line=1.12),
    T(104, 780, 740, 72, R("Equipo 13", 48, "white", b=True)),
    T(104, 872, 740, 64, LBL("Ideathon NEA Tech 2026", 40, "cobalt_light")),
    # 16:9 exacto sin recortar (928 x 522)
    I("00-portada.png", 896, 260, 928, 522, border=True),
], dark=True, chrome=False, notes=(
    "Somos el Equipo 13 y creamos Sobra: un mapa de oferta y demanda de biomasa para que los residuos de aserraderos "
    "y molinos arroceros de Corrientes encuentren comprador con un número que dice si el negocio cierra."))

# 2 · El Problema
slide([
    T(104, 96, 1200, 64, LBL("01 · El problema", 40, "cobalt")),
    T(104, 180, 1600, 130, R("Se quema lo que es combustible", 84, "fg", b=True)),
    B(104, 340, 540, 440, "surface", border_color="line"),
    T(144, 380, 460, 140, [R("200 mil", 84, "coral", b=True, mono=True)]),
    T(144, 540, 460, 210, R("toneladas al año se queman solo en Santa Rosa.", 42, "muted"), line=1.12),
    B(690, 340, 540, 440, "surface", border_color="line"),
    T(730, 380, 460, 140, [R("48,5 %", 84, "cobalt", b=True, mono=True)]),
    T(730, 540, 460, 210, R("del arroz del país sale de Corrientes.", 42, "muted"), line=1.12),
    B(1276, 340, 540, 440, "surface", border_color="line"),
    T(1316, 380, 460, 140, [R("516 mil", 84, "fg", b=True, mono=True)]),
    T(1316, 540, 460, 210, R("hectáreas forestadas: primera potencia nacional.", 42, "muted"), line=1.12),
    T(104, 816, 1440, 130, R("Falta saber quién compra y si el camión se come la ganancia.", 44, "fg", b=True)),
], notes=(
    "200.000 toneladas por año se queman solo en Santa Rosa, más la cascarilla del primer productor arrocero del país. "
    "No es basura: es combustible. Pero el que la tiene no sabe quién la compra ni si el flete se come la ganancia."))

# 3 · Por qué pasa (El flete)
slide([
    T(104, 96, 1200, 64, LBL("02 · El cuello de botella", 40, "cobalt")),
    T(104, 180, 1600, 130, R("Un camión lleno de aire", 84, "fg", b=True)),
    B(104, 340, 820, 540, "surface", border_color="line"),
    T(152, 376, 724, 64, LBL("Cascarilla suelta · 108 kg/m³", 40, "coral")),
    T(152, 460, 724, 150, R("25 km máx", 92, "coral", b=True, mono=True)),
    T(152, 632, 724, 220, R("70 % del camión viaja con aire: solo entran 9,7 t de las 30 t posibles.", 42, "muted"), line=1.12),
    B(992, 340, 820, 540, "surface", border_color="line"),
    T(1040, 376, 724, 64, LBL("En pellet · 650 kg/m³", 40, "green")),
    T(1040, 460, 724, 150, R("310 km (12x)", 92, "green", b=True, mono=True)),
    T(1040, 632, 724, 220, R("Carga las 30 t completas y multiplica 12 veces el radio rentable.", 42, "muted"), line=1.12),
], notes=(
    "La cascarilla suelta pesa tan poco que el camión se llena de aire antes que de carga: más allá de 25 km no conviene. "
    "En pellet entra toda la carga y viaja 12 veces más lejos, hasta 310 km."))

# 4 · La solución: Sobra.
slide([
    T(104, 96, 700, 64, LBL("03 · La solución", 40, "cobalt")),
    T(104, 180, 720, 240, [
        [R("¿Cierra el", 84, "fg", b=True)],
        [R("negocio?", 84, "cobalt", b=True)],
    ], line=1.0),
    T(104, 450, 720, 130, [R("01  ", 42, "cobalt", b=True, mono=True), R("13 plantas y 10 industrias reales", 42, "fg", b=True)]),
    T(104, 600, 720, 130, [R("02  ", 42, "cobalt", b=True, mono=True), R("Simula suelto o pellet en 1 clic", 42, "fg", b=True)]),
    T(104, 750, 720, 130, [R("03  ", 42, "cobalt", b=True, mono=True), R("Semáforo de margen neto real", 42, "fg", b=True)]),
    # 16:9 exacto sin recortar (960 x 540)
    I("01-inicio.png", 856, 260, 960, 540, border=True),
], notes=(
    "Por eso creamos Sobra: un mapa de oferta y demanda de biomasa que para cada vendedor y cada comprador "
    "responde una sola pregunta: ¿cierra el negocio?"))

# 5 · El momento clave: De Suelto a Pellet
slide([
    T(104, 96, 1500, 64, LBL("04 · Demo real · Molinos Ala a Frigorífico Duhalde (88 km)", 40, "cobalt")),
    T(104, 176, 1600, 115, R("De pérdida a ganancia en un clic", 76, "fg", b=True)),
    # Izquierda 16:9 (800 x 450)
    I("02-suelto.png", 104, 310, 800, 450, border=True),
    T(104, 780, 800, 60, LBL("Suelto · Radio 25 km · No cierra", 40, "coral")),
    T(104, 856, 800, 76, R("-$ 2,26 M / mes (-281 %)", 48, "coral", b=True, mono=True)),
    # Derecha 16:9 (800 x 450)
    I("03-pellet.png", 960, 310, 800, 450, border=True),
    T(960, 780, 800, 60, LBL("Pellet · Radio 310 km · Cierra", 40, "green")),
    T(960, 856, 600, 76, R("+$ 1,93 M / mes", 48, "green", b=True, mono=True)),
], notes=(
    "Molinos Ala en Mercedes quiere venderle al Frigorífico Duhalde a 88 km. Si va suelta, pierde 2,26 millones por mes. "
    "Con un clic a pellet, viaja 12 veces más lejos y el trato deja casi 2 millones de pesos por mes."))

# 6 · Ahorro para el comprador + Fuentes oficiales
slide([
    T(104, 96, 760, 64, LBL("05 · Ahorro y fuentes", 40, "cobalt")),
    B(104, 184, 720, 330, "cobalt_dark"),
    T(144, 212, 640, 120, R("-$ 15 M / mes", 76, "cobalt_light", b=True, mono=True)),
    T(144, 352, 640, 140, R("Ahorra el frigorífico frente al gas envasado (GLP).", 42, "white"), line=1.12),
    B(104, 544, 720, 360, "surface", border_color="line"),
    T(144, 576, 640, 64, LBL("Cada número con su fuente", 40, "cobalt")),
    T(144, 660, 640, 210, R("Tarifas FETRA 2026, GLP 2026, CONICET e IPCC con etiqueta Real o Estimado.", 42, "muted"), line=1.12),
    # Derecha 16:9 (960 x 540)
    I("07-fuentes.png", 856, 260, 960, 540, border=True),
], notes=(
    "Y al frigorífico le ahorra 15 millones por mes frente al gas envasado que utiliza hoy. "
    "Además, cada número de la plataforma muestra si es real o estimado y enlaza a su fuente oficial."))

# 7 · Mercado vivo + Modelo de negocio e Impacto
slide([
    T(104, 96, 1500, 64, LBL("06 · Mercado activo e impacto", 40, "cobalt")),
    T(104, 176, 800, 104, R("Publicás en 1 minuto", 68, "fg", b=True)),
    # Izquierda 16:9 (800 x 450)
    I("06-campanita.png", 104, 300, 800, 450, border=True),
    T(104, 774, 800, 140, R("Sin cuenta: contacto directo por WhatsApp y alertas de negocio.", 42, "muted"), line=1.12),
    # Derecha: 3 bloques limpios con alturas amplias
    B(960, 176, 856, 232, "surface", border_color="line"),
    T(1000, 198, 776, 60, LBL("1 · Freemium B2B + Comisión", 40, "cobalt")),
    T(1000, 276, 776, 120, R("Publicar gratis; plan pro y comisión por trato.", 40, "muted")),
    B(960, 432, 856, 232, "surface", border_color="line"),
    T(1000, 454, 776, 60, LBL("2 · Tablero para Corrientes", 40, "cobalt")),
    T(1000, 532, 776, 120, R("Menos humo, menos GLP y datos para planificar.", 40, "muted")),
    B(960, 688, 856, 224, "cobalt_dark"),
    T(1000, 708, 776, 90, R("~500 t CO2 / año", 60, "cobalt_light", b=True, mono=True)),
    T(1000, 816, 560, 76, R("Evitadas por acuerdo.", 42, "white")),
], notes=(
    "En Sobra publicás lo que tenés o buscás en un minuto sin crear cuenta, ves el contacto directo y recibís alertas. "
    "Para el aserradero es ingreso, para la industria energía barata y un solo acuerdo evita 500 toneladas de CO2 al año."))

# 8 · Cierre (Dark Cobalt)
slide([
    I("icon.png", 104, 104, 112, 112, border=False),
    T(240, 100, 560, 140, [R("Sobra", 96, "white", b=True), R(".", 96, "cobalt_light", b=True)]),
    T(104, 272, 1080, 260, [
        [R("Con un número", 88, "white", b=True)],
        [R("que te cierra.", 88, "cobalt_light", b=True)],
    ], line=1.0),
    T(104, 560, 1080, 100, R("ideathon-agro-web.vercel.app", 60, "cobalt_light", b=True, mono=True, link=DEMO)),
    T(104, 688, 1080, 76, R("Equipo 13 · Ideathon NEA Tech 2026", 48, "white", b=True)),
    T(104, 792, 1080, 150, R("Idea y diseño: Equipo 13. Programación con asistencia de IA. Datos reales con fuente oficial.", 40, "line"), line=1.15),
    I("qr-demo.png", 1312, 240, 480, 480, border=True, link=DEMO),
    T(1312, 756, 480, 72, LBL("Probá Sobra ahora", 40, "cobalt_light"), align="center"),
], dark=True, chrome=False, notes=(
    "Por esto los invitamos a probar Sobra, que pasa del residuo al mercado, con un número que te cierra."))


# ---------------- validación estricta de cajas y desborde ----------------
def estimate_text_height(e):
    total_h = 0
    for runs in e["paras"]:
        big = max(r["size"] for r in runs)
        text = "".join(r["t"] for r in runs)
        avg_char_w = big * (0.62 if any(r["mono"] for r in runs) else 0.56)
        chars_per_line = max(1, int(e["w"] / avg_char_w))
        lines = max(1, math.ceil(len(text) / chars_per_line))
        total_h += lines * big * LH * e["line"]
    return total_h


def validate():
    errs = []
    for n, sd in enumerate(SLIDES, 1):
        boxes = [e for e in sd["els"] if e["role"] != "bg"]
        for e in boxes:
            if e["x"] < SAFE or e["y"] < SAFE or e["x"] + e["w"] > W - SAFE or e["y"] + e["h"] > H - SAFE:
                errs.append(f"slide {n}: fuera del margen de seguridad: {label(e)}")
            if e["kind"] == "text":
                est_h = estimate_text_height(e)
                if est_h > e["h"] + 6:
                    errs.append(f"slide {n}: texto puede desbordar alto ({est_h:.0f}px > {e['h']}px): {label(e)}")
                if e["role"] == "content":
                    for p in e["paras"]:
                        for r in p:
                            if r["size"] < MIN_BODY_PX:
                                errs.append(f"slide {n}: texto menor a 20 pt: {r['t']}")
        content_boxes = [e for e in boxes if e["kind"] != "rect"]
        for i, a in enumerate(content_boxes):
            for b in content_boxes[i + 1:]:
                if (a["x"] < b["x"] + b["w"] + GAP and b["x"] < a["x"] + a["w"] + GAP and
                        a["y"] < b["y"] + b["h"] + GAP and b["y"] < a["y"] + a["h"] + GAP):
                    errs.append(f"slide {n}: se tocan {label(a)} y {label(b)}")
    if errs:
        raise SystemExit("Validación de cajas falló:\n  " + "\n  ".join(errs))


def label(e):
    return e.get("name") or e["paras"][0][0]["t"][:35]


# ---------------- PPTX ----------------
def E(v):
    return Emu(int(round(v * 6350)))


def rgb(name):
    return RGBColor.from_string(C.get(name, name))


ALIGN = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}
ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}


def crop_for(e):
    iw, ih = Image.open(IMG / e["name"]).size
    sc = max(e["w"] / iw, e["h"] / ih)
    dw, dh = iw * sc, ih * sc
    ex, ey = max(0.0, (dw - e["w"]) / dw), max(0.0, (dh - e["h"]) / dh)
    px, py = e["pos"]
    return px * ex, (1 - px) * ex, py * ey, (1 - py) * ey


def build_pptx():
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(W), E(H)
    for sd in SLIDES:
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = rgb("cobalt_dark" if sd.get("dark") else "canvas")
        for e in sd["els"]:
            if e["kind"] == "img":
                p = s.shapes.add_picture(str(IMG / e["name"]), E(e["x"]), E(e["y"]), E(e["w"]), E(e["h"]))
                p.crop_left, p.crop_right, p.crop_top, p.crop_bottom = crop_for(e)
                if e["border"]:
                    p.line.color.rgb = rgb("cobalt")
                    p.line.width = Pt(1.5)
                if e["link"]:
                    p.click_action.hyperlink.address = e["link"]
            elif e["kind"] == "rect":
                shp = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, E(e["x"]), E(e["y"]), E(e["w"]), E(e["h"]))
                shp.fill.solid()
                shp.fill.fore_color.rgb = rgb(e["fill"])
                if e["alpha"] < 1:
                    clr = shp.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
                    etree.SubElement(clr, qn("a:alpha")).set("val", str(int(e["alpha"] * 100000)))
                if e.get("border_color"):
                    shp.line.color.rgb = rgb(e["border_color"])
                    shp.line.width = Pt(1.0)
                else:
                    shp.line.fill.background()
                shp.shadow.inherit = False
            else:
                tb = s.shapes.add_textbox(E(e["x"]), E(e["y"]), E(e["w"]), E(e["h"]))
                tf = tb.text_frame
                tf.word_wrap = True
                tf.auto_size = MSO_AUTO_SIZE.NONE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                tf.vertical_anchor = ANCHOR[e["anchor"]]
                for i, runs in enumerate(e["paras"]):
                    para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
                    para.alignment = ALIGN[e["align"]]
                    para.line_spacing = e["line"]
                    for r in runs:
                        run = para.add_run()
                        run.text = r["t"]
                        f = run.font
                        f.name = MONO if r["mono"] else SANS
                        f.size = Pt(r["size"] * 0.5)
                        f.bold = r["b"]
                        f.color.rgb = rgb(r["color"])
                        if r["sp"]:
                            run._r.get_or_add_rPr().set("spc", str(int(round(r["sp"] * 0.5 * 100))))
                        if r["link"]:
                            run.hyperlink.address = r["link"]
        s.notes_slide.notes_text_frame.text = sd["notes"]
    prs.save(OUT)
    return len(prs.slides)


def build_html(path):
    imgs = IMG.as_uri()
    css = f"""
* {{ box-sizing: border-box; margin: 0; padding: 0; }}
body {{ background: #cbd5e1; }}
.slide {{ position: relative; width: {W}px; height: {H}px; overflow: hidden; margin-bottom: 20px; }}
.el {{ position: absolute; }}
.t {{ display: flex; flex-direction: column; overflow: visible; -webkit-font-smoothing: antialiased; }}
.t p {{ overflow-wrap: normal; }}
img.el {{ object-fit: cover; }}
"""
    out = [f"<!doctype html><html lang='es'><head><meta charset='utf-8'><title>PPTX Twin Check</title><style>{css}</style></head><body>"]
    for sd in SLIDES:
        bg = C["cobalt_dark"] if sd.get("dark") else C["canvas"]
        out.append(f"<section class='slide' style='background:#{bg}'>")
        for e in sd["els"]:
            box = f"left:{e['x']}px;top:{e['y']}px;width:{e['w']}px;height:{e['h']}px;"
            if e["kind"] == "img":
                px, py = e["pos"]
                border = f"border:2px solid #{C['cobalt']};" if e["border"] else ""
                out.append(f"<img class='el' data-role='{e['role']}' src='{imgs}/{e['name']}' style='{box}{border}object-position:{px * 100}% {py * 100}%'>")
            elif e["kind"] == "rect":
                c = C[e["fill"]]
                bdr = f"border:1px solid #{C[e['border_color']]};" if e.get("border_color") else ""
                out.append(f"<div class='el' data-role='bg' style='{box}background:#{c};{bdr}'></div>")
            else:
                just = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}[e["anchor"]]
                out.append(f"<div class='el t' data-role='{e['role']}' style='{box}justify-content:{just};text-align:{e['align']}'>")
                for runs in e["paras"]:
                    big = max(r["size"] for r in runs)
                    lh = big * LH * e["line"]
                    out.append(f"<p style='line-height:{lh:.1f}px;font-size:{big}px'>")
                    for r in runs:
                        fam = f'"{MONO}", monospace' if r["mono"] else f'"{SANS}", Arial, sans-serif'
                        st = (f"font-family:{fam};font-size:{r['size']}px;color:#{C[r['color']]};"
                              f"font-weight:{700 if r['b'] else 400};letter-spacing:{r['sp']}px")
                        out.append(f"<span style='{st}'>{r['t']}</span>")
                    out.append("</p>")
                out.append("</div>")
        out.append("</section>")
    out.append("</body></html>")
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text("".join(out), encoding="utf-8")


validate()
n = build_pptx()
print(f"PPTX: {OUT} | {n} slides | validación de cajas y alturas OK")
if HTML_OUT:
    build_html(HTML_OUT)
    print(f"HTML: {HTML_OUT}")
